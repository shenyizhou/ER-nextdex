from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def normalize_text(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def set_if_present(target, source, translated):
    if source and translated:
        target[source] = translated


def build_glossary(workbook_path: Path):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    glossary = {
        "exact": {},
        "species": {},
        "abilities": {},
        "abilityDescriptions": {},
        "moves": {},
        "moveDescriptions": {},
        "items": {},
        "maps": {},
        "trainerClasses": {},
        "trainers": {},
        "regex": [],
    }

    species_sheet = workbook["简易图鉴"]
    for row in species_sheet.iter_rows(min_row=3, values_only=True):
        english_name = normalize_text(row[5] if len(row) > 5 else "")
        chinese_name = normalize_text(row[6] if len(row) > 6 else "")
        set_if_present(glossary["species"], english_name, chinese_name)

    ability_sheet = workbook["特性"]
    for row in ability_sheet.iter_rows(min_row=2, values_only=True):
        english_name = normalize_text(row[2] if len(row) > 2 else "")
        chinese_name = normalize_text(row[3] if len(row) > 3 else "")
        english_desc = normalize_text(row[4] if len(row) > 4 else "")
        chinese_desc = normalize_text(row[5] if len(row) > 5 else "")
        set_if_present(glossary["abilities"], english_name, chinese_name)
        set_if_present(glossary["abilityDescriptions"], english_name, chinese_desc)
        set_if_present(glossary["exact"], english_desc, chinese_desc)

    moves_sheet = workbook["招式"]
    for row in moves_sheet.iter_rows(min_row=2, values_only=True):
        english_name = normalize_text(row[2] if len(row) > 2 else "")
        chinese_name = normalize_text(row[3] if len(row) > 3 else "")
        english_desc = normalize_text(row[9] if len(row) > 9 else "")
        chinese_desc = normalize_text(row[10] if len(row) > 10 else "")
        set_if_present(glossary["moves"], english_name, chinese_name)
        set_if_present(glossary["moveDescriptions"], english_name, chinese_desc)
        set_if_present(glossary["exact"], english_desc, chinese_desc)

    return glossary


def main():
    workbook_path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("ER2.5正式版图鉴v0.7.xlsm")
    output_path = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("src/translate_gamedata.glossary.json")
    glossary = build_glossary(workbook_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(glossary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    summary = {
        "workbook": str(workbook_path),
        "output": str(output_path),
        "species": len(glossary["species"]),
        "abilities": len(glossary["abilities"]),
        "abilityDescriptions": len(glossary["abilityDescriptions"]),
        "moves": len(glossary["moves"]),
        "moveDescriptions": len(glossary["moveDescriptions"]),
        "exact": len(glossary["exact"]),
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
